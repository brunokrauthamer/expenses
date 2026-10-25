"""Importa despesas das planilhas mensais de cartão.

Mostra uma prévia por padrão. Use --apply para gravar no banco.
Execute após alembic upgrade head.
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable
from uuid import UUID

from openpyxl import load_workbook

from app.modules.expenses.import_contracts import ImportExpenseCommand
from app.modules.expenses.import_repository import ExpenseImportRepository
from app.modules.expenses.import_service import ExpenseImportService
from app.shared.exceptions import DomainValidationError
from app.shared.sqlalchemy_uow import SqlAlchemyUnitOfWork

MONTHS = {
    "jan": 1,
    "fev": 2,
    "mar": 3,
    "abr": 4,
    "mai": 5,
    "jun": 6,
    "jul": 7,
    "ago": 8,
    "set": 9,
    "out": 10,
    "nov": 11,
    "dez": 12,
}
MONTH_IN_NAME = re.compile(
    r"(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)_(\d{4})",
    re.IGNORECASE,
)
INSTALLMENT = re.compile(r"^(\d+)\s*/\s*(\d+)$")
CENT = Decimal("0.01")
CATEGORY_ALIASES = {
    "aluguel + cond": "Moradia",
    "presente": "Presentes",
    "itens para casa": "Itens para casa",
    "vícios": "Lazer",
    "lazer mandi": "Lazer",
    "vestuário": "Lazer",
    "outro": "Mercado",
}
IGNORED_CATEGORIES = {"empréstimo"}


@dataclass(frozen=True)
class SourceRow:
    path: Path
    row_number: int
    expense_date: date
    installment_number: int | None
    installment_count: int | None
    payment_method: str
    category: str
    description: str
    amount: Decimal

    @property
    def source_key(self) -> str:
        return f"{self.path.name}|Página1!{self.row_number}"

    @property
    def grouping_key(self) -> tuple[str, str, int]:
        return (
            _normalize(self.description),
            _normalize(_payment_method(self.payment_method) or self.payment_method),
            self.installment_count or 0,
        )


@dataclass(frozen=True)
class PreparedExpense:
    description: str
    category_name: str
    payment_method_name: str
    expense_date: date
    amounts: tuple[Decimal, ...]
    sources: tuple[SourceRow, ...]
    note: str = ""

    @property
    def total(self) -> Decimal:
        return sum(self.amounts, Decimal("0.00"))


def _normalize(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.strip().casefold())
    return " ".join("".join(char for char in value if not unicodedata.combining(char)).split())


def _payment_method(value: str) -> str | None:
    value = _normalize(value)
    if value == "btg":
        return "Cartão de Crédito"
    if value.startswith("pix"):
        return "PIX"
    return None


def _category(value: str, existing: Iterable[str]) -> str | None:
    aliases = {_normalize(alias): target for alias, target in CATEGORY_ALIASES.items()}
    normalized = _normalize(value)
    if normalized in aliases:
        return aliases[normalized]
    return next((name for name in existing if _normalize(name) == normalized), None)


def _month(path: Path) -> date:
    match = MONTH_IN_NAME.search(path.name)
    if match is None:
        raise ValueError(f"Nome de planilha sem mês/ano reconhecido: {path.name}")
    return date(int(match.group(2)), MONTHS[match.group(1).lower()], 1)


def _amount(value: object) -> Decimal:
    try:
        return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Valor inválido: {value!r}") from error


def _read(paths: list[Path]):
    rows: list[SourceRow] = []
    issues: list[str] = []
    subscriptions = nonpositive = ignored_category_count = 0
    for path in paths:
        month_date = _month(path)
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            sheet = workbook.worksheets[0]
            for number, values in enumerate(
                sheet.iter_rows(min_row=2, max_col=5, values_only=True), start=2
            ):
                if not any(value is not None for value in values):
                    continue
                parcel, payment, category, description, raw_amount = values
                if str(parcel or "").strip().casefold() == "assinatura":
                    subscriptions += 1
                    continue
                if _normalize(str(category or "")) in {
                    _normalize(value) for value in IGNORED_CATEGORIES
                }:
                    ignored_category_count += 1
                    continue
                try:
                    amount = _amount(raw_amount)
                except ValueError as error:
                    issues.append(f"{path.name}, linha {number}: {error}")
                    continue
                if amount <= 0:
                    nonpositive += 1
                    continue
                if not all(str(value or "").strip() for value in (payment, category, description)):
                    issues.append(f"{path.name}, linha {number}: faltam campos obrigatórios.")
                    continue
                token = str(parcel or "").strip()
                if not token:
                    installment_number = installment_count = None
                else:
                    match = INSTALLMENT.fullmatch(token)
                    if not match:
                        issues.append(
                            f"{path.name}, linha {number}: parcela não reconhecida ({token!r})."
                        )
                        continue
                    installment_number, installment_count = map(int, match.groups())
                    if installment_number < 1 or installment_count < installment_number:
                        issues.append(f"{path.name}, linha {number}: parcela inválida ({token!r}).")
                        continue
                rows.append(
                    SourceRow(
                        path,
                        number,
                        month_date,
                        installment_number,
                        installment_count,
                        str(payment).strip(),
                        str(category).strip(),
                        str(description).strip(),
                        amount,
                    )
                )
        finally:
            workbook.close()
    return rows, issues, subscriptions, nonpositive, ignored_category_count


def _row_order(row: SourceRow):
    return row.amount, row.row_number


def _month_offset(day: date, offset: int) -> date:
    index = day.year * 12 + day.month - 1 + offset
    year, month_zero = divmod(index, 12)
    return date(year, month_zero + 1, 1)


def _prepare(rows: list[SourceRow], category_names: list[str], oldest_month: date):
    regular: dict[tuple[str, str, int], list[SourceRow]] = defaultdict(list)
    for row in rows:
        if row.installment_number is not None:
            regular[row.grouping_key].append(row)

    prepared: list[PreparedExpense] = []
    issues: list[str] = []
    estimates = ignored_later = 0
    consumed: set[str] = set()

    def add(row, amounts, sources, description=None, note=""):
        category = _category(row.category, category_names)
        payment = _payment_method(row.payment_method)
        if category is None:
            issues.append(
                f"Categoria sem mapeamento: {row.category!r} ({row.path.name}, linha {row.row_number})."
            )
            return
        if payment is None:
            issues.append(
                f"Meio sem mapeamento: {row.payment_method!r} ({row.path.name}, linha {row.row_number})."
            )
            return
        label = description or row.description
        if len(label) > 120:
            issues.append(
                f"Descrição acima de 120 caracteres: {label!r} ({row.path.name}, linha {row.row_number})."
            )
            return
        prepared.append(
            PreparedExpense(
                label, category, payment, row.expense_date, tuple(amounts), tuple(sources), note
            )
        )
        consumed.update(source.source_key for source in sources)

    for row in rows:
        if row.installment_number is None:
            add(row, (row.amount,), (row,))

    for group_rows in regular.values():
        by_month: dict[date, list[SourceRow]] = defaultdict(list)
        for row in group_rows:
            by_month[row.expense_date].append(row)
        months = sorted(by_month)
        first_month = oldest_month
        legacy = [
            row
            for row in by_month.get(first_month, [])
            if row.installment_number is not None and row.installment_number > 1
        ]
        legacy_rows = list(legacy)
        active = legacy
        for month in months[1:]:
            following = []
            current = by_month[month]
            expected: dict[int, list[SourceRow]] = defaultdict(list)
            for previous in active:
                expected[previous.installment_number + 1].append(previous)  # type: ignore[operator]
            for number, previous_rows in expected.items():
                candidates = sorted(
                    [
                        item
                        for item in current
                        if item.installment_number == number
                        and item.source_key not in consumed
                        and item not in legacy_rows
                    ],
                    key=_row_order,
                )
                previous_rows = sorted(previous_rows, key=_row_order)
                following.extend(candidates[: min(len(candidates), len(previous_rows))])
            legacy_rows.extend(following)
            active = following

        for row in legacy_rows:
            add(
                row,
                (row.amount,),
                (row,),
                description=f"{row.description}, parcela {row.installment_number}",
            )

        first_rows = [row for row in group_rows if row.installment_number == 1]
        for row in sorted(first_rows, key=lambda item: (item.expense_date, _row_order(item))):
            next_month = _month_offset(row.expense_date, 1)
            candidates = sorted(
                [
                    item
                    for item in by_month.get(next_month, [])
                    if item.installment_number == 2 and item.source_key not in consumed
                ],
                key=_row_order,
            )
            same_month_firsts = sorted(
                [
                    item
                    for item in first_rows
                    if item.expense_date == row.expense_date and item.source_key not in consumed
                ],
                key=_row_order,
            )
            index = same_month_firsts.index(row)
            second = candidates[index] if index < len(candidates) else None
            if second:
                amounts = (row.amount,) + (second.amount,) * (row.installment_count - 1)  # type: ignore[operator]
                sources = (row, second)
                note = ""
            else:
                amounts = (row.amount,) * row.installment_count  # type: ignore[arg-type]
                sources = (row,)
                note = "Estimado pela primeira parcela × número de parcelas"
                estimates += 1
            add(row, amounts, sources, note=note)

        ignored_later += sum(
            1
            for row in group_rows
            if row.installment_number is not None
            and row.installment_number >= 3
            and row.source_key not in consumed
        )
        for row in group_rows:
            if row.installment_number == 2 and row.source_key not in consumed:
                issues.append(
                    f"Parcela 2 sem primeira parcela correspondente: {row.description!r} "
                    f"({row.path.name}, linha {row.row_number})."
                )
    return prepared, issues, estimates, ignored_later


def _paths(arguments: list[str]) -> list[Path]:
    if arguments:
        result = [Path(item).expanduser().resolve() for item in arguments]
    else:
        result = [
            item.resolve()
            for item in (Path.home() / "Downloads").glob("*_2026.xlsx")
            if "gastos cartao" in _normalize(item.name)
        ]
    result.sort(key=_month)
    if not result:
        raise ValueError("Nenhuma planilha mensal encontrada.")
    dates = [_month(path) for path in result]
    if len(set(dates)) != len(dates):
        raise ValueError("Há mais de uma planilha para o mesmo mês.")
    missing = [str(path) for path in result if not path.is_file()]
    if missing:
        raise ValueError("Arquivos não encontrados: " + ", ".join(missing))
    return result


def _preview(
    paths,
    prepared,
    issues,
    subscriptions,
    nonpositive,
    ignored_categories,
    estimates,
    ignored_later,
):
    print(f"Planilhas: {len(paths)} | {paths[0].name} a {paths[-1].name}")
    print(
        f"Compras: {len(prepared)} | Total: R$ {sum((item.total for item in prepared), Decimal('0.00')):.2f}"
    )
    print(f"Assinaturas ignoradas: {subscriptions}")
    print(f"Valores negativos/zero ignorados: {nonpositive}")
    print(f"Categoria Empréstimo ignorada: {ignored_categories}")
    print(f"Parcelas posteriores ignoradas: {ignored_later}")
    print(f"Total estimado pela primeira parcela: {estimates}")
    print(f"Pendências: {len(issues)}")
    for issue in issues:
        print(f"- {issue}")
    print("\nPrévia das compras:")
    for item in prepared:
        label = f"{item.expense_date:%Y-%m} | {item.description} | {item.category_name}"
        label += f" | {item.payment_method_name} | R$ {item.total:.2f}"
        if len(item.amounts) > 1:
            label += f" | {len(item.amounts)} cobranças"
        if item.note:
            label += f" | {item.note}"
        print(label)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("files", nargs="*", help="Planilhas mensais .xlsx")
    parser.add_argument("--apply", action="store_true", help="Grava após mostrar a prévia.")
    parser.add_argument(
        "--user-id", required=True, type=UUID, help="UUID da conta que receberá as despesas"
    )
    args = parser.parse_args()
    paths = _paths(args.files)
    rows, read_issues, subscriptions, nonpositive, ignored_categories = _read(paths)
    from app.database import SessionLocal

    with SessionLocal() as database:
        service = ExpenseImportService(
            ExpenseImportRepository(database),
            SqlAlchemyUnitOfWork(database),
        )
        try:
            categories = service.user_categories(args.user_id)
        except DomainValidationError as error:
            raise SystemExit(str(error)) from error
        oldest_month = min(_month(path) for path in paths)
        prepared, prepare_issues, estimates, ignored_later = _prepare(
            rows, categories, oldest_month
        )
        issues = read_issues + prepare_issues
        _preview(
            paths,
            prepared,
            issues,
            subscriptions,
            nonpositive,
            ignored_categories,
            estimates,
            ignored_later,
        )
        if not args.apply:
            print("\nPrévia apenas; o banco não foi alterado.")
            return
        if issues:
            raise SystemExit("Importação cancelada. Resolva as pendências listadas.")
        commands = [
            ImportExpenseCommand(
                description=item.description,
                category_name=item.category_name,
                payment_method_name=item.payment_method_name,
                expense_date=item.expense_date,
                amounts=item.amounts,
                source_keys=tuple(source.source_key for source in item.sources),
            )
            for item in prepared
        ]
        try:
            count = service.apply(args.user_id, commands)
        except DomainValidationError as error:
            raise SystemExit(str(error)) from error
        if count == 0 and commands:
            print("Todas as linhas já foram importadas. Nenhuma alteração feita.")
        print(f"\nImportação concluída: {count} compras.")


if __name__ == "__main__":
    main()
